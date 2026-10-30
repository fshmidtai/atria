"""Tender report — сборка клиентского Excel под кворк «База тендеров 44-ФЗ».

Первый заказ → python infra/deliver/tender_report.py [archive.jsonl] [out.xlsx]
→ файл уходит покупателю. Никаких ручных выгрузок: данные уже лежат в архиве
колонии (storage/tenderslice/archive.jsonl), кворк обещает ровно то, что рой
уже собрал.

Колонки: № · Закупка · Сумма ₽ · Регион · Приём до · Заказчик · Ссылка
Фильтры (по умолчанию как в объявлении): сумма ≤ 300 000 ₽, только с суммой.
Сортировка: ближайший дедлайн первый (закрытые — в конец).
"""
from __future__ import annotations

import json
import re
import sys
from datetime import datetime, timezone
from pathlib import Path

LIMIT = 300_000
COLS = ["№", "Закупка", "Сумма ₽", "Регион", "Приём заявок до",
        "Заказчик", "Ссылка на тендер"]


def parse_sum(v) -> int | None:
    if v is None:
        return None
    m = re.findall(r"[\d,]+", str(v).replace(" ", ""))
    if not m:
        return None
    try:
        return int(m[-1].replace(",", ""))
    except ValueError:
        return None


def parse_closes(v) -> datetime | None:
    if not v:
        return None
    for fmt in ("%d %b %Y, %H:%M %Z", "%d %b %Y, %H:%M"):
        try:
            dt = datetime.strptime(str(v).strip(), fmt)
            return dt.replace(tzinfo=timezone.utc) if dt.tzinfo is None else dt
        except ValueError:
            continue
    return None


def rows(archive: list[dict], limit: int = LIMIT,
         max_rows: int = 200) -> list[list]:
    out = []
    now = datetime.now(timezone.utc)
    for it in archive:
        s = parse_sum(it.get("sum"))
        if s is None or s > limit:
            continue
        closes = parse_closes(it.get("closes"))
        out.append({
            "title": str(it.get("title", ""))[:200],
            "sum": s,
            "region": str(it.get("region", "")).split(",")[0].strip(),
            "closes": closes,
            "buyer": str(it.get("buyer", "")).strip(),
            "url": str(it.get("url", "")),
        })
    # открытые ближе к дедлайну — первыми; закрытые в конец
    out.sort(key=lambda r: (r["closes"] is None or r["closes"] < now,
                            r["closes"] or now, -r["sum"]))
    out = out[:max_rows]
    return [[i + 1, r["title"], r["sum"], r["region"],
             (r["closes"].strftime("%d %b %Y, %H:%M UTC") if r["closes"]
              else (str(it.get("closes", "")) or "—")),
             r["buyer"] or "—", r["url"]]
            for i, r in enumerate(out)]


def to_xlsx(rows: list[list], out: Path, title: str = "База тендеров 44-ФЗ"):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Тендеры"
    ws.append(COLS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3864")
        c.alignment = Alignment(vertical="center")
    for r in rows:
        ws.append(r)
    widths = [5, 62, 14, 22, 24, 30, 60]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:G{len(rows) + 1}"
    ws.append([])
    ws.append(["Собрано автономным роем Атрия · freetender.net (публичные "
               "данные) · фильтр: сумма ≤ 300 000 ₽"])
    wb.save(out)


def to_csv(rows: list[list]) -> str:
    import csv
    import io
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(COLS)
    w.writerows(rows)
    return buf.getvalue()


def main() -> int:
    src = Path(sys.argv[1] if len(sys.argv) > 1
               else "storage/tenderslice/archive.jsonl")
    if not src.exists():
        print(f"archive not found: {src}")
        return 1
    archive = [json.loads(l) for l in src.read_text(encoding="utf-8").splitlines()
               if l.strip()]
    data = rows(archive)
    print(f"archive={len(archive)} → отфильтровано ≤{LIMIT}: {len(data)}")
    dst = Path(sys.argv[2] if len(sys.argv) > 2
               else src.parent / "tender_report.xlsx")
    if str(dst).endswith(".csv"):
        dst.write_text(to_csv(data), encoding="utf-8")
    else:
        to_xlsx(data, dst)
    print("saved:", dst)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
